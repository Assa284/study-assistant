import io
import json
import os
import re
import requests
import time
import unicodedata
import streamlit as st
import streamlit.components.v1 as components
from docx import Document
from google import genai
from google.genai import types
from openpyxl import load_workbook
from html import escape as html_escape

st.set_page_config(
    page_title="Study Assistant", page_icon="📚", initial_sidebar_state="expanded"
)

# Keep what the student typed or chose when they move between pages.
for _k in (
    "language", "task", "n_series", "notes", "save_history", "pref_language", "pref_task",
    "pref_series", "pref_rate", "pref_download", "pref_textsize",
):
    if _k in st.session_state:
        st.session_state[_k] = st.session_state[_k]

APP_CSS = """
<style>
html { font-size: __FONT__%; }
.stApp {
  background-image:
    radial-gradient(900px 500px at 90% -5%, rgba(124, 92, 255, 0.20), transparent 60%),
    radial-gradient(700px 420px at -5% 105%, rgba(34, 211, 238, 0.12), transparent 55%);
}
[data-testid="stHeader"] { background: transparent; }
[data-testid="stSidebar"] { border-right: 1px solid rgba(128, 128, 128, 0.25); }
div[data-testid="stVerticalBlockBorderWrapper"] { border-radius: 18px; }
.stButton > button, .stDownloadButton > button, .stLinkButton > a { border-radius: 12px; }
button[kind="primary"], button[data-testid="stBaseButton-primary"] {
  background: linear-gradient(90deg, #22d3ee, #7c5cff);
  border: 0; font-weight: 700;
}
button[kind="primary"] p, button[data-testid="stBaseButton-primary"] p { color: #06101f !important; }
textarea, input { border-radius: 12px !important; }
h1 { font-weight: 800; letter-spacing: -0.5px; }
.avatar {
  width: 38px; height: 38px; border-radius: 50%; flex: none;
  background: linear-gradient(135deg, #7c5cff, #22d3ee); color: #06101f; font-weight: 800;
  display: flex; align-items: center; justify-content: center;
}
.acct { display: flex; align-items: center; gap: 10px; margin: 6px 0 10px; }
.acct-name { font-weight: 600; word-break: break-word; }
</style>
"""
_scale = {"Small": 90, "Medium": 100, "Large": 115}.get(
    st.session_state.get("pref_textsize", "Medium"), 100
)
st.markdown(APP_CSS.replace("__FONT__", str(_scale)), unsafe_allow_html=True)

def get_secret(name):
    value = os.environ.get(name)
    if value:
        return value
    try:
        return st.secrets[name]
    except Exception:
        return None


api_key = get_secret("GEMINI_API_KEY")
if not api_key:
    st.error("API key not found. Set GEMINI_API_KEY first.")
    st.stop()

access_code = get_secret("ACCESS_CODE")
if access_code and not st.session_state.get("unlocked"):
    st.title("📚 Study Assistant")
    entered = st.text_input("Access code", type="password")
    if entered == access_code:
        st.session_state["unlocked"] = True
        st.rerun()
    st.info("Enter your access code to use the app.")
    st.stop()

client = genai.Client(api_key=api_key)
MODELS = [
    "gemini-3.5-flash-lite",
    "gemini-3.1-flash-lite",
    "gemini-3.8-flash",
    "gemini-3.5-flash",
]
# For PDFs and photos (tables, handwriting) the stronger models are tried first;
# if one is busy or out of free quota, the app falls back to the next automatically.
MODELS_STRONG = [
    "gemini-3.8-flash",
    "gemini-3.5-flash",
    "gemini-3.5-flash-lite",
    "gemini-3.1-flash-lite",
]
MAX_MB = 10
MAX_SOURCES = 5

TASKS = ["Summary", "Full explanation", "Quiz", "Explain simply", "Ask a question"]

STYLE = (
    "Start directly with the content: no greeting, no introduction and no closing remarks. "
    "Write EVERYTHING, including headings, in the answer language. "
    "Do the requested task: never just copy the material or translate it back. "
    "Use the correct standard terms of the subject in the answer language; only for key "
    "technical terms (a few words each, never whole sentences or lines), add the original "
    "term in square brackets [like this] the first time it appears, but only when it is different from the "
    "word you wrote; never repeat the same word in brackets. "
    "Translate the meaning, not word for word: write natural, clear sentences like a "
    "good textbook. Copy technical terms of the original language exactly as they are "
    "spelled in the material. "
    "Use clean Markdown. Never write notes, corrections, comments about the quality of the "
    "text, or your own reasoning. "
    "Copy numbers, names and currency symbols exactly as written; if a word or number is "
    "unclear or missing (for example in a photo or a table), leave that detail out instead "
    "of guessing or leaving a blank such as 'is .'. Never assume or invent a number "
    "that you did not clearly read; write 'not clear in the document' instead. "
    "Keep abbreviations exactly as the material writes them (for example CIR). "
    "Translate section titles such as 'Series', 'Answer key' and 'Difficult words' into "
    "the answer language (in French: 'Série', 'Corrigé', 'Mots difficiles'). "
    "If the material asks a question, do not turn it into a statement and do not invent "
    "its answer. "
    "For calculations and long tables inside a document (for example worked examples in "
    "a PDF), never paste jumbled numbers: explain what they calculate and the steps in words; quote a figure "
    "only if you are sure it is exactly right. "
)


def instruction_for(task, n):
    if task == "Summary":
        return (
            "Summarize the material in at most 10 clear bullet points of one or two lines each "
            "(up to 15 if the material is long). Keep only the key ideas, and include the "
            "main worked example in one bullet if there is one. The summary must be much "
            "shorter than the material."
        )
    if task == "Full explanation":
        return (
            "Explain the material in detail, point by point, in the order it appears. "
            "For every point give a clear explanation and, when useful, a short example. "
            "If the material has a worked example with calculations, walk through it step "
            "by step in words: what is given, what is calculated, in which order, and what "
            "the result means. Write a figure only if you can read it clearly and it agrees "
            "with the other figures; otherwise describe the step without the number. "
            "At the end add a section called 'Difficult words' with a short, simple meaning "
            "(one line each) for every hard or technical word of the material: at least 8 "
            "words if the material is long."
        )
    if task == "Quiz":
        return (
            f"Create {n} different quiz series about the material, titled 'Series 1', "
            f"'Series 2' and so on up to 'Series {n}'. Each series has 5 multiple-choice "
            "questions with 4 options (A-D). Together the series must cover ALL the topics "
            "of the material. Never repeat or reword a question from another series. "
            "Base every question and answer only on what the material says, and check each "
            "answer key against the material. "
            "Put ONE answer key at the end of each series (never after each question): a numbered list with the correct letter and one short line of explanation; never explain the wrong options. "
            "Format: a heading '## Series 1' for each series; each question on its own "
            "line in bold; its four options as a bullet list under it; then 'Answer key' "
            "as a numbered list. If the material is too short to make "
            f"{n} series without repeating a topic, make fewer series and say so in one line."
        )
    if task == "Explain simply":
        return (
            "Explain the material in very simple words, as if to a secondary school "
            "student. Rewrite everything in much simpler words and do not reuse the sentences of the material. Use one everyday example. Keep it under 250 words, with no tables and no "
            "headings copied from the material."
        )
    return (
        "Answer the student's question clearly and accurately in a few short paragraphs, "
        "using the material if it is provided, and say which document your answer comes from. "
        "If the material does not contain the answer, "
        "say so in one clear sentence, then give a short general answer starting with "
        "'From general knowledge (not from your documents):' (translated into the answer "
        "language). If you are not sure about a "
        "fact, say so instead of guessing. If the question is too vague to answer well, ask "
        "the student one short clarifying question instead of guessing."
    )


def read_docx(f):
    doc = Document(io.BytesIO(f.getvalue()))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def read_xlsx(f):
    wb = load_workbook(io.BytesIO(f.getvalue()), data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def ask_ai(contents, strong=False):
    last_error = None
    for model in (MODELS_STRONG if strong else MODELS):
        for attempt in range(2):
            try:
                response = client.models.generate_content(
                    model=model, contents=contents
                )
                return response, model
            except Exception as e:
                text = str(e)
                last_error = e
                if "503" in text or "UNAVAILABLE" in text:
                    if attempt == 0:
                        time.sleep(4)
                        continue
                    break
                if (
                    "404" in text
                    or "NOT_FOUND" in text
                    or "429" in text
                    or "RESOURCE_EXHAUSTED" in text
                ):
                    break
                print("AI error:", last_error)
                raise RuntimeError(
                    "The AI could not process this. Try again, or use a clearer or smaller file."
                )
    print("AI error:", last_error)
    raise RuntimeError(
        "The AI is busy or the free limit is reached right now. Please try again in a few minutes."
    )


def collect_sources(files, notes, include_notes, camera=None):
    """Return one source per uploaded file (and one for typed text)."""
    sources, problems = [], []
    for f in files or []:
        name = f.name.lower()
        if f.size > MAX_MB * 1024 * 1024:
            problems.append(f"{f.name} is bigger than {MAX_MB} MB and was skipped.")
            continue
        try:
            if name.endswith(".docx"):
                sources.append({"label": f.name, "text": read_docx(f), "media": []})
            elif name.endswith(".xlsx"):
                sources.append({"label": f.name, "text": read_xlsx(f), "media": []})
            elif name.endswith(".pdf"):
                part = types.Part.from_bytes(
                    data=f.getvalue(), mime_type="application/pdf"
                )
                sources.append({"label": f.name, "text": "", "media": [part]})
            else:
                mime = "image/png" if name.endswith(".png") else "image/jpeg"
                part = types.Part.from_bytes(data=f.getvalue(), mime_type=mime)
                sources.append({"label": f.name, "text": "", "media": [part]})
        except Exception:
            problems.append(f"Could not read {f.name}.")
    if camera is not None:
        part = types.Part.from_bytes(
            data=camera.getvalue(), mime_type=camera.type or "image/jpeg"
        )
        sources.append({"label": "Camera photo", "text": "", "media": [part]})
    if include_notes and notes.strip():
        sources.append({"label": "Your typed text", "text": notes, "media": []})
    return sources, problems


def add_runs(paragraph, text):
    """Write text into a Word paragraph, turning **bold** into real bold."""
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part.replace("*", ""))


def build_docx(items):
    """items = list of (title, markdown_text). Returns the bytes of a Word file."""
    doc = Document()
    for n, (title, md) in enumerate(items):
        if n:
            doc.add_page_break()
        doc.add_heading(title, level=1)
        for raw in md.split("\n"):
            line = raw.rstrip()
            if not line.strip():
                continue
            heading = re.match(r"^(#{1,6})\s+(.*)", line)
            bullet = re.match(r"^(\s*)[-*\u2022]\s+(.*)", line)
            if heading:
                level = min(len(heading.group(1)) + 1, 4)
                doc.add_heading(heading.group(2).replace("**", ""), level=level)
            elif bullet:
                style = "List Bullet 2" if len(bullet.group(1)) >= 2 else "List Bullet"
                add_runs(doc.add_paragraph(style=style), bullet.group(2))
            else:
                add_runs(doc.add_paragraph(), line.strip())
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def safe_name(text):
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9_-]+", "_", text).strip("_")[:40] or "result"


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


PDF_CHARS = {"→": "->", "←": "<-", "≥": ">=", "≤": "<=", "≠": "!=", "≈": "~", "−": "-"}


def pdf_safe(text):
    """The standard PDF fonts only know Western characters; swap or drop the rest."""
    for a, b in PDF_CHARS.items():
        text = text.replace(a, b)
    return text.replace("$", "").encode("cp1252", "ignore").decode("cp1252")


def build_pdf(items):
    """items = list of (title, markdown_text). Returns PDF bytes, or None if unavailable."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate
    except ImportError:
        return None

    def markup(t):
        t = pdf_safe(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
        return t.replace("*", "")

    try:
        base = getSampleStyleSheet()
        body = ParagraphStyle("body", parent=base["BodyText"], fontSize=10.5, leading=14, spaceAfter=4)
        bullet1 = ParagraphStyle("b1", parent=body, leftIndent=16, bulletIndent=4)
        bullet2 = ParagraphStyle("b2", parent=body, leftIndent=32, bulletIndent=20)
        story = []
        for n, (title, md) in enumerate(items):
            if n:
                story.append(PageBreak())
            story.append(Paragraph(markup(title), base["Heading1"]))
            for raw in md.split("\n"):
                line = raw.rstrip()
                if not line.strip():
                    continue
                heading = re.match(r"^(#{1,6})\s+(.*)", line)
                bullet = re.match(r"^(\s*)[-*\u2022]\s+(.*)", line)
                if heading:
                    style = base["Heading2"] if len(heading.group(1)) <= 2 else base["Heading3"]
                    story.append(Paragraph(markup(heading.group(2).replace("**", "")), style))
                elif bullet:
                    style = bullet2 if len(bullet.group(1)) >= 2 else bullet1
                    story.append(Paragraph(markup(bullet.group(2)), style, bulletText="\u2022"))
                else:
                    story.append(Paragraph(markup(line.strip()), body))
        buffer = io.BytesIO()
        SimpleDocTemplate(
            buffer, pagesize=A4, leftMargin=50, rightMargin=50, topMargin=50, bottomMargin=50
        ).build(story)
        return buffer.getvalue()
    except Exception:
        return None


def download_buttons(items, base_name, key):
    """Word and PDF buttons side by side. The student can choose which ones to see in Settings."""
    mode = st.session_state.get("pref_download", "Word and PDF")
    c1, c2 = st.columns(2)
    if mode != "PDF only":
        with c1:
            st.download_button(
                "⬇️ Word", data=build_docx(items), file_name=f"{base_name}.docx",
                mime=DOCX_MIME, key=f"{key}_docx",
            )
    if mode != "Word only":
        with (c1 if mode == "PDF only" else c2):
            pdf = build_pdf(items)
            if pdf:
                st.download_button(
                    "⬇️ PDF", data=pdf, file_name=f"{base_name}.pdf",
                    mime="application/pdf", key=f"{key}_pdf",
                )
            else:
                st.caption("PDF is not available yet.")


def speech_text(md):
    """Turn the answer into clean text for the voice: no symbols, no [original terms]."""
    text = re.sub(r"\[[^\]]*\]", "", md)
    text = re.sub(r"[\U00010000-\U0010ffff]", "", text)  # emojis
    text = re.sub(r"[#*_`>|$\\]", " ", text)
    lines = [re.sub(r"^[-\u2022]\s+", "", ln.strip()) for ln in text.split("\n")]
    lines = [ln for ln in lines if ln]
    return " ".join(ln if ln[-1] in ".!?:;" else ln + "." for ln in lines)


SPEAK_HTML = """
<div style="font-family:sans-serif">
  <div style="display:flex;gap:8px">
    <button id="play" style="padding:6px 12px;border-radius:8px;border:1px solid #ccc;background:#fff;cursor:pointer">🔊 Listen</button>
    <button id="stop" style="padding:6px 12px;border-radius:8px;border:1px solid #ccc;background:#fff;cursor:pointer">⏹ Stop</button>
  </div>
  <div id="info" style="font-size:12px;color:#666;margin-top:4px"></div>
</div>
<script>
const text = __TEXT__;
const lang = "__LANG__";
const wanted = "__NAME__";
const prefix = lang.slice(0, 2).toLowerCase();
const synth = window.speechSynthesis;
const info = document.getElementById("info");

function loadVoices() {
  return new Promise(resolve => {
    const now = synth.getVoices();
    if (now.length) return resolve(now);
    synth.onvoiceschanged = () => resolve(synth.getVoices());
    setTimeout(() => resolve(synth.getVoices()), 1500);
  });
}

document.getElementById("stop").onclick = () => synth.cancel();
document.getElementById("play").onclick = async () => {
  synth.cancel();
  const voices = await loadVoices();
  const norm = v => v.lang.toLowerCase().replace("_", "-");
  const voice = voices.find(v => norm(v) === lang.toLowerCase())
             || voices.find(v => norm(v).startsWith(prefix));
  if (voices.length && !voice) {
    info.textContent = "No " + wanted + " voice found on this device. Use the Online voice button below, or install a voice in the text-to-speech settings.";
    return;
  }
  info.textContent = voice ? "Voice: " + voice.name + " (" + voice.lang + ")" : "";
  const parts = text.match(/[^.!?]+[.!?:;]*/g) || [text];
  parts.forEach(p => {
    if (!p.trim()) return;
    const u = new SpeechSynthesisUtterance(p.trim());
    u.lang = lang;
    u.rate = __RATE__;
    if (voice) u.voice = voice;
    synth.speak(u);
  });
};
</script>
"""


def speak_widget(md, lang_choice):
    """Read a result aloud in the language it was written in (not the current menu choice)."""
    french = lang_choice == "Français"
    html = (
        SPEAK_HTML.replace("__TEXT__", json.dumps(speech_text(md)))
        .replace("__LANG__", "fr-FR" if french else "en-US")
        .replace("__NAME__", "French" if french else "English")
        .replace("__RATE__", str(float(st.session_state.get("pref_rate", 1.0))))
    )
    components.html(html, height=70)


ONLINE_VOICE_MAX = 2500  # characters read aloud by the online voice


@st.cache_data(show_spinner=False, max_entries=50)
def make_online_voice(text, lang):
    """Make an mp3 on the server, so it works even if the device has no voice."""
    from gtts import gTTS

    buffer = io.BytesIO()
    gTTS(text=text, lang=lang).write_to_fp(buffer)
    return buffer.getvalue()


def online_voice(md, lang_choice, key):
    spoken = speech_text(md)
    if st.button("🌐 Online voice (any device)", key=f"{key}_btn"):
        with st.spinner("Preparing the voice..."):
            try:
                st.session_state[key] = make_online_voice(
                    spoken[:ONLINE_VOICE_MAX], "fr" if lang_choice == "Français" else "en"
                )
            except ImportError:
                st.caption("The online voice is not available yet.")
            except Exception:
                st.caption("The online voice is busy right now. Please try again in a minute.")
    if st.session_state.get(key):
        st.audio(st.session_state[key], format="audio/mp3")
        if len(spoken) > ONLINE_VOICE_MAX:
            st.caption("Only the first part is read aloud.")


def run_one(label, text, media):
    source_rule = (
        "" if task == "Ask a question"
        else "Only use information from the provided material and do not add facts of "
        "your own (everyday examples are allowed). "
    )
    prompt = (
        f"You are a friendly study assistant for students. ANSWER LANGUAGE: {language} only. "
        f"{instruction_for(task, n_series)} {source_rule}{STYLE}\n\n"
        f"MATERIAL (from: {label}):\n{text}\n\n"
        f"REMINDER: write your whole answer in {language}, even if the material is in another language."
    )
    response, used_model = ask_ai([*media, prompt], strong=bool(media))
    answer = getattr(response, "text", None)
    if not answer or not answer.strip():
        raise RuntimeError("The AI returned an empty answer. Please try again.")
    try:
        cut_off = "MAX_TOKENS" in str(response.candidates[0].finish_reason)
    except Exception:
        cut_off = False
    if cut_off:
        answer += (
            "\n\n⚠️ This answer was cut off because it was too long. "
            "Try fewer quiz series or a shorter file."
        )
    return answer, used_model


def is_signed_in():
    return bool(getattr(st.user, "is_logged_in", False))


def user_email():
    return (st.user.get("email") or "").strip().lower()


def history_ready():
    """History works only when the student is signed in and the database secrets exist."""
    return bool(
        get_secret("SUPABASE_URL") and get_secret("SUPABASE_KEY")
        and is_signed_in() and user_email()
    )


def sb_headers(extra=None):
    key = get_secret("SUPABASE_KEY")
    headers = {"apikey": key, "Content-Type": "application/json"}
    if key.startswith("eyJ"):  # old-style keys also need the Authorization header
        headers["Authorization"] = f"Bearer {key}"
    if extra:
        headers.update(extra)
    return headers


def sb_url():
    return get_secret("SUPABASE_URL").rstrip("/") + "/rest/v1/history"


def db_reason(exc):
    """A short, safe explanation of why a database call failed (never includes keys)."""
    resp = getattr(exc, "response", None)
    if resp is not None:
        try:
            msg = resp.json().get("message") or resp.text
        except Exception:
            msg = resp.text
        return f"HTTP {resp.status_code}: {str(msg)[:120]}"
    return type(exc).__name__


@st.cache_data(ttl=120, show_spinner=False)
def fetch_history(email):
    resp = requests.get(
        sb_url(),
        headers=sb_headers(),
        params={
            "user_email": f"eq.{email}",
            "select": "id,created_at,task,language,source_label,result_text",
            "order": "created_at.desc",
            "limit": "100",
        },
        timeout=15,
    )
    resp.raise_for_status()
    return resp.json()


def save_history(email, results):
    rows = [
        {
            "user_email": email,
            "task": r.get("task", ""),
            "language": r.get("lang", ""),
            "source_label": r["label"],
            "result_text": r["text"],
            "model": r["model"],
        }
        for r in results
        if r["model"]  # skip error messages
    ]
    if rows:
        resp = requests.post(
            sb_url(),
            headers=sb_headers({"Prefer": "return=minimal"}),
            data=json.dumps(rows),
            timeout=15,
        )
        resp.raise_for_status()
        fetch_history.clear()


def delete_history(email, row_id=None):
    params = {"user_email": f"eq.{email}"}  # always limited to the signed-in student
    if row_id is not None:
        params["id"] = f"eq.{row_id}"
    resp = requests.delete(sb_url(), headers=sb_headers(), params=params, timeout=15)
    resp.raise_for_status()
    fetch_history.clear()
    st.session_state.pop("history_pick", None)


FORM_URL = "https://docs.google.com/forms/d/e/1FAIpQLSd2Ye-U8eMSdyUrBcAU5eMIUvGYhRUPPMwejPOOLcFnH1unIA/viewform"


def account_block():
    """Sidebar account area. Shows the sign-in button only when the [auth] secrets exist."""
    try:
        _ = st.secrets["auth"]
    except Exception:
        st.caption("Guest mode")
        return
    if is_signed_in():
        name = st.user.get("name") or st.user.get("email") or "Student"
        initials = "".join(w[0] for w in name.split()[:2]).upper() or "S"
        st.markdown(
            f'<div class="acct"><div class="avatar">{html_escape(initials)}</div>'
            f'<div class="acct-name">{html_escape(name)}</div></div>',
            unsafe_allow_html=True,
        )
        st.button("Log out", on_click=st.logout, key="logout_btn")
    else:
        st.button("👤 Sign in with Google", on_click=st.login, key="login_btn")
        st.caption("Sign in to save your work.")


def apply_defaults():
    """Copy the default choices from Settings to the Home controls."""
    st.session_state["language"] = st.session_state.get("pref_language", "English")
    st.session_state["task"] = st.session_state.get("pref_task", "Summary")
    st.session_state["n_series"] = st.session_state.get("pref_series", 3)


def reset_settings():
    for k in ("pref_language", "pref_task", "pref_series", "pref_rate", "pref_download",
              "pref_textsize", "save_history"):
        st.session_state.pop(k, None)
    for k in ("language", "task", "n_series"):
        st.session_state.pop(k, None)


def history_page():
    st.title("🕘 History")
    if not is_signed_in():
        st.info("Sign in with Google (in the left menu) to save your work and see it here.")
        return
    if not history_ready():
        st.info("Saved history is not set up yet.")
        return
    email = user_email()
    try:
        rows = fetch_history(email)
    except Exception as e:
        print("History load failed:", db_reason(e))
        st.caption(f"Your history is not available right now. ({db_reason(e)})")
        return
    if not rows:
        st.caption("Nothing saved yet. Your next results will appear here.")
        return
    with st.container(border=True):
        pick = st.selectbox(
            "Open a saved result",
            range(len(rows)),
            format_func=lambda k: f"{rows[k]['created_at'][:10]} · {rows[k]['task']} · {rows[k]['source_label']}",
            key="history_pick",
        )
        row = rows[pick]
        st.markdown(row["result_text"])
        title = f"{row['source_label']} - {row['task']}"
        download_buttons(
            [(title, row["result_text"])],
            f"{safe_name(row['source_label'])}_{safe_name(row['task'])}",
            f"hist_{row['id']}",
        )
        if st.button("🗑 Delete this result", key=f"del_{row['id']}"):
            deleted = False
            try:
                delete_history(email, row["id"])
                deleted = True
            except Exception:
                st.caption("Could not delete right now. Please try again.")
            if deleted:
                st.rerun()


def settings_page():
    st.title("⚙️ Settings")
    st.session_state.setdefault("pref_language", "English")
    st.session_state.setdefault("pref_task", "Summary")
    st.session_state.setdefault("pref_series", 3)
    st.session_state.setdefault("pref_rate", 1.0)
    st.session_state.setdefault("pref_download", "Word and PDF")
    st.session_state.setdefault("pref_textsize", "Medium")

    with st.container(border=True):
        st.subheader("Defaults")
        st.radio("Default language", ["English", "Français"], horizontal=True,
                 key="pref_language", on_change=apply_defaults)
        st.selectbox("Default task", TASKS, key="pref_task", on_change=apply_defaults)
        st.slider("Default number of quiz series", 1, 5, key="pref_series", on_change=apply_defaults)

    with st.container(border=True):
        st.subheader("Reading aloud")
        st.slider("Reading speed", 0.6, 1.4, step=0.1, key="pref_rate")
        st.caption("Used by the 🔊 Listen button. The online voice has a fixed speed.")

    with st.container(border=True):
        st.subheader("Downloads")
        st.radio("Show download buttons for", ["Word and PDF", "Word only", "PDF only"],
                 horizontal=True, key="pref_download")

    with st.container(border=True):
        st.subheader("Appearance")
        st.radio("Text size", ["Small", "Medium", "Large"], horizontal=True, key="pref_textsize")
        st.caption("For dark or light mode, open the ⋮ menu at the top right, choose Settings, then Theme.")

    with st.container(border=True):
        st.subheader("Privacy and data")
        if history_ready():
            st.checkbox("💾 Save my work to my history", value=True, key="save_history")
            st.caption(
                "Results are saved with your Google email and only you can see them. "
                "Your uploaded files and photos are not saved."
            )
            if st.checkbox("I want to delete my whole history", key="del_all_ok") and st.button(
                "🗑 Delete all my history", key="del_all"
            ):
                try:
                    delete_history(user_email())
                    st.success("Your history was deleted.")
                except Exception:
                    st.caption("Could not delete right now. Please try again.")
        else:
            st.caption("Sign in with Google to save your work and manage your history.")
        st.caption("The 🌐 online voice sends the text of a result to Google's voice service.")

    with st.container(border=True):
        st.subheader("Help and about")
        st.markdown(
            "1. Paste notes, upload files, or take a 📷 photo.\n"
            "2. Pick a language and what you want.\n"
            "3. Press **Get started**, then listen or download the result."
        )
        st.link_button("💬 Send us your feedback", FORM_URL)
        st.caption("Study Assistant · AI can make mistakes. Check important facts in your textbook.")

    st.button("↩️ Reset my settings", on_click=reset_settings)


with st.sidebar:
    st.markdown("## 📚 Study Assistant")
    page = st.radio(
        "Menu", ["🏠 Home", "🕘 History", "⚙️ Settings"],
        label_visibility="collapsed", key="nav",
    )
    st.divider()
    account_block()

if page == "🕘 History":
    history_page()
    st.stop()
if page == "⚙️ Settings":
    settings_page()
    st.stop()

# ---------------- Home ----------------
st.session_state.setdefault("language", st.session_state.get("pref_language", "English"))
st.session_state.setdefault("task", st.session_state.get("pref_task", "Summary"))
st.session_state.setdefault("n_series", st.session_state.get("pref_series", 3))

head_left, head_right = st.columns([3, 1])
with head_left:
    st.title("📚 Study Assistant")
    st.caption(
        "Your personal study companion. Upload, paste or write your notes, "
        "choose what you want, and get instant help."
    )
with head_right:
    with st.popover("📷 Photo"):
        camera = st.camera_input("Take a photo of your notes")
if camera is not None:
    st.caption("📷 A camera photo is ready. It will be used when you press **Get started**.")

with st.container(border=True):
    language = st.radio("Language / Langue", ["English", "Français"], horizontal=True, key="language")
    task = st.selectbox("What do you want?", TASKS, key="task")
    n_series = 3
    if task == "Quiz":
        n_series = st.slider("How many quiz series?", 1, 5, key="n_series")
    notes = st.text_area("Paste your notes (or type your question) here", height=200, key="notes")
    voice = None
    if task == "Ask a question" and hasattr(st, "audio_input"):
        voice = st.audio_input("🎤 Or ask by voice (record, then press Get started)")
    files = st.file_uploader(
        "Or upload files: PDF, Word (.docx), Excel (.xlsx), or a photo of your notes",
        type=["pdf", "docx", "xlsx", "png", "jpg", "jpeg"],
        accept_multiple_files=True,
    )
    go = st.button("Get started →", type="primary")

if "run_id" not in st.session_state:
    st.session_state["run_id"] = 0

if go:
    results = []
    if task == "Ask a question":
        sources, problems = collect_sources(files, notes, include_notes=False, camera=camera)
        for p in problems:
            st.warning(p)
        question = notes.strip()
        if not question and voice is not None:
            with st.spinner("Listening..."):
                try:
                    heard, _ = ask_ai(
                        [
                            types.Part.from_bytes(
                                data=voice.getvalue(), mime_type="audio/wav"
                            ),
                            "Write down exactly what the student says, in the language "
                            "spoken. Return only the text, nothing else.",
                        ]
                    )
                    question = (getattr(heard, "text", "") or "").strip()
                except Exception:
                    st.error(
                        "Could not understand the recording. Try again or type your question."
                    )
        if not question:
            st.warning("Please type or record your question first.")
        else:
            text = "\n\n".join(
                f"[Document: {s['label']}]\n{s['text']}" for s in sources if s["text"]
            )
            text += f"\n\nSTUDENT QUESTION:\n{question}"
            media = [m for s in sources for m in s["media"]]
            with st.spinner("Working on it..."):
                try:
                    answer, used_model = run_one("your documents and question", text, media)
                    if not notes.strip():
                        answer = f"**🎤 You asked:** {question}\n\n{answer}"
                    results.append(
                        {"label": "Answer", "text": answer, "model": used_model, "task": task, "lang": language}
                    )
                except Exception as e:
                    st.error(str(e))
    elif not files and camera is None and notes.strip().endswith("?") and len(notes.strip()) < 300:
        st.info(
            "It looks like you typed a question. Choose **Ask a question** in the menu "
            "above to get an answer."
        )
    else:
        sources, problems = collect_sources(files, notes, include_notes=True, camera=camera)
        for p in problems:
            st.warning(p)
        if not sources:
            st.warning("Please paste some notes or upload a file first.")
        else:
            if len(sources) > MAX_SOURCES:
                st.warning(
                    f"Only the first {MAX_SOURCES} items were used. Please do the rest in a second round."
                )
                sources = sources[:MAX_SOURCES]
            for i, s in enumerate(sources, start=1):
                with st.spinner(f"Working on {s['label']} ({i}/{len(sources)})..."):
                    try:
                        answer, used_model = run_one(s["label"], s["text"], s["media"])
                        results.append(
                            {"label": s["label"], "text": answer, "model": used_model, "task": task, "lang": language}
                        )
                    except Exception as e:
                        results.append(
                            {"label": s["label"], "text": f"⚠️ {e}", "model": None, "task": task}
                        )
    if results:
        st.session_state["results"] = results
        st.session_state["run_id"] += 1
        if history_ready() and st.session_state.get("save_history", True):
            try:
                save_history(user_email(), results)
                st.caption("💾 Saved to your history.")
            except Exception as e:
                print("History save failed:", db_reason(e))
                st.caption(f"Could not save to your history right now. ({db_reason(e)})")
    else:
        st.session_state.pop("results", None)

# Show the last results (kept in memory so they stay on screen after a thumbs click)
if "results" in st.session_state:
    models_used = {r["model"] for r in st.session_state["results"] if r["model"]}
    run_id = st.session_state["run_id"]
    items = []
    for i, r in enumerate(st.session_state["results"]):
        st.subheader(f"📄 {r['label']}")
        st.markdown(r["text"])
        if r["model"]:  # no listen/download for error messages
            task_name = r.get("task", "result")
            title = f"{r['label']} - {task_name}"
            items.append((title, r["text"]))
            speak_widget(r["text"], r.get("lang", language))
            online_voice(r["text"], r.get("lang", language), f"voice_{run_id}_{i}")
            download_buttons(
                [(title, r["text"])],
                f"{safe_name(r['label'])}_{safe_name(task_name)}",
                f"dl_{run_id}_{i}",
            )
        st.divider()
    if len(items) > 1:
        st.markdown("**All results together**")
        download_buttons(items, "Study_Assistant_results", f"dl_all_{run_id}")
    st.caption(
        "AI can make mistakes. Check important facts in your textbook. "
        f"(Model: {', '.join(sorted(models_used)) or 'none'})"
    )

    st.write("**Was this helpful?**")
    rating = st.feedback("thumbs", key=f"thumbs_{st.session_state['run_id']}")
    if rating is not None:
        st.success(
            "Thanks for your feedback! \U0001F64F"
            if rating == 1
            else "Sorry about that. Please tell us what went wrong in the form below."
        )

    FORM_URL = "https://docs.google.com/forms/d/e/1FAIpQLSd2Ye-U8eMSdyUrBcAU5eMIUvGYhRUPPMwejPOOLcFnH1unIA/viewform"
    st.link_button("\U0001F4AC Give us your feedback (30 seconds)", FORM_URL)
