import io
import os
import time
import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from openpyxl import load_workbook

st.set_page_config(page_title="Study Assistant", page_icon="📚")
st.title("📚 Study Assistant")
st.write("Paste your notes or upload a file, choose what you want, and pick a language.")

def get_secret(name):
    value = os.environ.get(name)
    if value:
        return value
    try:
        return st.secrets[name]
    except Exception:
        return None


api_key = get_secret("GEMINI_API_KEY")
if not api_key:
    st.error("API key not found. Set GEMINI_API_KEY first.")
    st.stop()

access_code = get_secret("ACCESS_CODE")
if access_code:
    entered = st.text_input("Access code", type="password")
    if entered != access_code:
        st.info("Enter your access code to use the app.")
        st.stop()

client = genai.Client(api_key=api_key)
MODELS = [
    "gemini-3.5-flash-lite",
    "gemini-3.1-flash-lite",
    "gemini-3.8-flash",
    "gemini-3.5-flash",
]
MAX_MB = 10

language = st.radio("Language / Langue", ["English", "Français"], horizontal=True)
task = st.selectbox(
    "What do you want?", ["Summary", "Quiz", "Explain simply", "Ask a question"]
)
notes = st.text_area("Paste your notes (or type your question) here", height=200)
files = st.file_uploader(
    "Or upload files: PDF, Word (.docx), Excel (.xlsx), or a photo of your notes",
    type=["pdf", "docx", "xlsx", "png", "jpg", "jpeg"],
    accept_multiple_files=True,
)

INSTRUCTIONS = {
    "Summary": "Summarize the material in clear bullet points. Keep only the key ideas.",
    "Quiz": "Create 5 multiple-choice questions from the material. Give 4 options (A-D) for each, then list the correct answers with a one-line explanation at the end.",
    "Explain simply": "Explain the material in very simple words, as if to a secondary school student. Use one everyday example.",
    "Ask a question": "Answer the student's question clearly and accurately in a few short paragraphs. If you are not sure about a fact, say so instead of guessing.",
}


def read_docx(f):
    doc = Document(io.BytesIO(f.getvalue()))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def read_xlsx(f):
    wb = load_workbook(io.BytesIO(f.getvalue()), data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def ask_ai(contents):
    last_error = None
    for model in MODELS:
        for attempt in range(2):
            try:
                response = client.models.generate_content(
                    model=model, contents=contents
                )
                return response, model
            except Exception as e:
                text = str(e)
                last_error = e
                if "503" in text or "UNAVAILABLE" in text:
                    if attempt == 0:
                        time.sleep(4)
                        continue
                    break
                if (
                    "404" in text
                    or "NOT_FOUND" in text
                    or "429" in text
                    or "RESOURCE_EXHAUSTED" in text
                ):
                    break
                raise
    raise RuntimeError(
        "All AI models are busy or out of free quota right now. "
        f"Please try again later. (Last error: {last_error})"
    )


if "run_id" not in st.session_state:
    st.session_state["run_id"] = 0

if st.button("Go"):
    text_parts = []
    media_parts = []
    problems = []

    for f in files or []:
        name = f.name.lower()
        if f.size > MAX_MB * 1024 * 1024:
            problems.append(f"{f.name} is bigger than {MAX_MB} MB and was skipped.")
            continue
        try:
            if name.endswith(".docx"):
                text_parts.append(f"[File: {f.name}]\n{read_docx(f)}")
            elif name.endswith(".xlsx"):
                text_parts.append(f"[File: {f.name}]\n{read_xlsx(f)}")
            elif name.endswith(".pdf"):
                media_parts.append(
                    types.Part.from_bytes(data=f.getvalue(), mime_type="application/pdf")
                )
            else:
                mime = "image/png" if name.endswith(".png") else "image/jpeg"
                media_parts.append(
                    types.Part.from_bytes(data=f.getvalue(), mime_type=mime)
                )
        except Exception:
            problems.append(f"Could not read {f.name}.")

    for p in problems:
        st.warning(p)

    if notes.strip():
        text_parts.append(notes)

    if not text_parts and not media_parts:
        st.warning("Please paste some notes or upload a file first.")
    else:
        if task == "Ask a question":
            source_rule = ""
        else:
            source_rule = "Only use information from the provided notes and files. "
        prompt = (
            f"You are a friendly study assistant for students. "
            f"{INSTRUCTIONS[task]} Write your whole answer in {language}. "
            f"{source_rule}\n\nMATERIAL:\n" + "\n\n".join(text_parts)
        )
        with st.spinner("Working on it..."):
            try:
                response, used_model = ask_ai([*media_parts, prompt])
                st.session_state["result"] = (response.text, used_model)
                st.session_state["run_id"] += 1
            except Exception as e:
                st.session_state.pop("result", None)
                st.error(f"Something went wrong: {e}")

# Show the last result (kept in memory so it stays on screen after a thumbs click)
if "result" in st.session_state:
    result_text, used_model = st.session_state["result"]
    st.markdown(result_text)
    st.caption(
        f"AI can make mistakes. Check important facts in your textbook. (Model: {used_model})"
    )

    st.divider()
    st.write("**Was this helpful?**")
    rating = st.feedback("thumbs", key=f"thumbs_{st.session_state['run_id']}")
    if rating is not None:
        st.success(
            "Thanks for your feedback! \U0001F64F"
            if rating == 1
            else "Sorry about that. Please tell us what went wrong in the form below."
        )

    FORM_URL = "https://docs.google.com/forms/d/e/1FAIpQLSd2Ye-U8eMSdyUrBcAU5eMIUvGYhRUPPMwejPOOLcFnH1unIA/viewform"
    st.link_button("\U0001F4AC Give us your feedback (30 seconds)", FORM_URL)
