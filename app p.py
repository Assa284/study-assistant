import io
import os
import time
import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from openpyxl import load_workbook

st.set_page_config(page_title="Study Assistant", page_icon="📚")
st.title("📚 Study Assistant")
st.write("Paste your notes or upload files, choose what you want, and pick a language.")


def get_secret(name):
    value = os.environ.get(name)
    if value:
        return value
    try:
        return st.secrets[name]
    except Exception:
        return None


api_key = get_secret("GEMINI_API_KEY")
if not api_key:
    st.error("API key not found. Set GEMINI_API_KEY first.")
    st.stop()

access_code = get_secret("ACCESS_CODE")
if access_code:
    entered = st.text_input("Access code", type="password")
    if entered != access_code:
        st.info("Enter your access code to use the app.")
        st.stop()

client = genai.Client(api_key=api_key)
MODELS = [
    "gemini-3.5-flash-lite",
    "gemini-3.1-flash-lite",
    "gemini-3.8-flash",
    "gemini-3.5-flash",
]
# For PDFs and photos (tables, handwriting) the stronger models are tried first;
# if one is busy or out of free quota, the app falls back to the next automatically.
MODELS_STRONG = [
    "gemini-3.8-flash",
    "gemini-3.5-flash",
    "gemini-3.5-flash-lite",
    "gemini-3.1-flash-lite",
]
MAX_MB = 10
MAX_SOURCES = 5

language = st.radio("Language / Langue", ["English", "Français"], horizontal=True)
task = st.selectbox(
    "What do you want?",
    ["Summary", "Full explanation", "Quiz", "Explain simply", "Ask a question"],
)
n_series = 3
if task == "Quiz":
    n_series = st.slider("How many quiz series?", 1, 5, 3)
notes = st.text_area("Paste your notes (or type your question) here", height=200)
files = st.file_uploader(
    "Or upload files: PDF, Word (.docx), Excel (.xlsx), or a photo of your notes",
    type=["pdf", "docx", "xlsx", "png", "jpg", "jpeg"],
    accept_multiple_files=True,
)

STYLE = (
    "Start directly with the content: no greeting, no introduction and no closing remarks. "
    "Write EVERYTHING, including headings, in the answer language. "
    "Do the requested task: never just copy the material or translate it back. "
    "Use the correct standard terms of the subject in the answer language; only for key "
    "technical terms (a few words each, never whole sentences or lines), add the original "
    "term in square brackets [like this] the first time it appears, but only when it is different from the "
    "word you wrote; never repeat the same word in brackets. "
    "Translate the meaning, not word for word: write natural, clear sentences like a "
    "good textbook. Copy technical terms of the original language exactly as they are "
    "spelled in the material. "
    "Use clean Markdown. Never write notes, corrections, comments about the quality of the "
    "text, or your own reasoning. "
    "Copy numbers, names and currency symbols exactly as written; if a word or number is "
    "unclear or missing (for example in a photo or a table), leave that detail out instead "
    "of guessing or leaving a blank such as 'is .'. "
    "If the material asks a question, do not turn it into a statement and do not invent "
    "its answer. "
    "For calculations and long tables inside a document (for example worked examples in "
    "a PDF), never paste jumbled numbers: explain what they calculate and the steps in words; quote a figure "
    "only if you are sure it is exactly right. "
)


def instruction_for(task, n):
    if task == "Summary":
        return (
            "Summarize the material in at most 10 clear bullet points of one or two lines each "
            "(up to 15 if the material is long). Keep only the key ideas, and include the "
            "main worked example in one bullet if there is one. The summary must be much "
            "shorter than the material."
        )
    if task == "Full explanation":
        return (
            "Explain the material in detail, point by point, in the order it appears. "
            "For every point give a clear explanation and, when useful, a short example. "
            "If the material has a worked example with calculations, walk through it step "
            "by step in words: what is given, what is calculated, in which order, and what "
            "the result means. Write a figure only if you can read it clearly and it agrees "
            "with the other figures; otherwise describe the step without the number. "
            "At the end add a section called 'Difficult words' with a short, simple meaning "
            "(one line each) for every hard or technical word of the material: at least 8 "
            "words if the material is long."
        )
    if task == "Quiz":
        return (
            f"Create {n} different quiz series about the material, titled 'Series 1', "
            f"'Series 2' and so on up to 'Series {n}'. Each series has 5 multiple-choice "
            "questions with 4 options (A-D). Together the series must cover ALL the topics "
            "of the material. Never repeat or reword a question from another series. "
            "Base every question and answer only on what the material says, and check each "
            "answer key against the material. "
            "Put ONE answer key at the end of each series (never after each question): a numbered list with the correct letter and one short line of explanation; never explain the wrong options. "
            "Format: a heading '## Series 1' for each series; each question on its own "
            "line in bold; its four options as a bullet list under it; then 'Answer key' "
            "as a numbered list. If the material is too short to make "
            f"{n} series without repeating a topic, make fewer series and say so in one line."
        )
    if task == "Explain simply":
        return (
            "Explain the material in very simple words, as if to a secondary school "
            "student. Rewrite everything in much simpler words and do not reuse the sentences of the material. Use one everyday example. Keep it under 250 words, with no tables and no "
            "headings copied from the material."
        )
    return (
        "Answer the student's question clearly and accurately in a few short paragraphs, "
        "using the material if it is provided, and say which document your answer comes from. "
        "If the material does not contain the answer, "
        "say so in one clear sentence, then give a short general answer starting with "
        "'From general knowledge (not from your documents):'. If you are not sure about a "
        "fact, say so instead of guessing. If the question is too vague to answer well, ask "
        "the student one short clarifying question instead of guessing."
    )


def read_docx(f):
    doc = Document(io.BytesIO(f.getvalue()))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def read_xlsx(f):
    wb = load_workbook(io.BytesIO(f.getvalue()), data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def ask_ai(contents, strong=False):
    last_error = None
    for model in (MODELS_STRONG if strong else MODELS):
        for attempt in range(2):
            try:
                response = client.models.generate_content(
                    model=model, contents=contents
                )
                return response, model
            except Exception as e:
                text = str(e)
                last_error = e
                if "503" in text or "UNAVAILABLE" in text:
                    if attempt == 0:
                        time.sleep(4)
                        continue
                    break
                if (
                    "404" in text
                    or "NOT_FOUND" in text
                    or "429" in text
                    or "RESOURCE_EXHAUSTED" in text
                ):
                    break
                print("AI error:", last_error)
                raise RuntimeError(
                    "The AI could not process this. Try again, or use a clearer or smaller file."
                )
    print("AI error:", last_error)
    raise RuntimeError(
        "The AI is busy or the free limit is reached right now. Please try again in a few minutes."
    )


def collect_sources(files, notes, include_notes):
    """Return one source per uploaded file (and one for typed text)."""
    sources, problems = [], []
    for f in files or []:
        name = f.name.lower()
        if f.size > MAX_MB * 1024 * 1024:
            problems.append(f"{f.name} is bigger than {MAX_MB} MB and was skipped.")
            continue
        try:
            if name.endswith(".docx"):
                sources.append({"label": f.name, "text": read_docx(f), "media": []})
            elif name.endswith(".xlsx"):
                sources.append({"label": f.name, "text": read_xlsx(f), "media": []})
            elif name.endswith(".pdf"):
                part = types.Part.from_bytes(
                    data=f.getvalue(), mime_type="application/pdf"
                )
                sources.append({"label": f.name, "text": "", "media": [part]})
            else:
                mime = "image/png" if name.endswith(".png") else "image/jpeg"
                part = types.Part.from_bytes(data=f.getvalue(), mime_type=mime)
                sources.append({"label": f.name, "text": "", "media": [part]})
        except Exception:
            problems.append(f"Could not read {f.name}.")
    if include_notes and notes.strip():
        sources.append({"label": "Your typed text", "text": notes, "media": []})
    return sources, problems


def run_one(label, text, media):
    source_rule = (
        "" if task == "Ask a question"
        else "Only use information from the provided material. "
    )
    prompt = (
        f"You are a friendly study assistant for students. ANSWER LANGUAGE: {language} only. "
        f"{instruction_for(task, n_series)} {source_rule}{STYLE}\n\n"
        f"MATERIAL (from: {label}):\n{text}\n\n"
        f"REMINDER: write your whole answer in {language}, even if the material is in another language."
    )
    response, used_model = ask_ai([*media, prompt], strong=bool(media))
    answer = getattr(response, "text", None)
    if not answer or not answer.strip():
        raise RuntimeError("The AI returned an empty answer. Please try again.")
    try:
        cut_off = "MAX_TOKENS" in str(response.candidates[0].finish_reason)
    except Exception:
        cut_off = False
    if cut_off:
        answer += (
            "\n\n⚠️ This answer was cut off because it was too long. "
            "Try fewer quiz series or a shorter file."
        )
    return answer, used_model


if "run_id" not in st.session_state:
    st.session_state["run_id"] = 0

if st.button("Go"):
    results = []
    if task == "Ask a question":
        sources, problems = collect_sources(files, notes, include_notes=False)
        for p in problems:
            st.warning(p)
        if not notes.strip():
            st.warning("Please type your question in the box first.")
        else:
            text = "\n\n".join(
                f"[Document: {s['label']}]\n{s['text']}" for s in sources if s["text"]
            )
            text += f"\n\nSTUDENT QUESTION:\n{notes}"
            media = [m for s in sources for m in s["media"]]
            with st.spinner("Working on it..."):
                try:
                    answer, used_model = run_one("your documents and question", text, media)
                    results.append({"label": "Answer", "text": answer, "model": used_model})
                except Exception as e:
                    st.error(str(e))
    elif not files and notes.strip().endswith("?") and len(notes.strip()) < 300:
        st.info(
            "It looks like you typed a question. Choose **Ask a question** in the menu "
            "above to get an answer."
        )
    else:
        sources, problems = collect_sources(files, notes, include_notes=True)
        for p in problems:
            st.warning(p)
        if not sources:
            st.warning("Please paste some notes or upload a file first.")
        else:
            if len(sources) > MAX_SOURCES:
                st.warning(
                    f"Only the first {MAX_SOURCES} items were used. Please do the rest in a second round."
                )
                sources = sources[:MAX_SOURCES]
            for i, s in enumerate(sources, start=1):
                with st.spinner(f"Working on {s['label']} ({i}/{len(sources)})..."):
                    try:
                        answer, used_model = run_one(s["label"], s["text"], s["media"])
                        results.append(
                            {"label": s["label"], "text": answer, "model": used_model}
                        )
                    except Exception as e:
                        results.append(
                            {"label": s["label"], "text": f"⚠️ {e}", "model": None}
                        )
    if results:
        st.session_state["results"] = results
        st.session_state["run_id"] += 1
    else:
        st.session_state.pop("results", None)

# Show the last results (kept in memory so they stay on screen after a thumbs click)
if "results" in st.session_state:
    models_used = {r["model"] for r in st.session_state["results"] if r["model"]}
    for r in st.session_state["results"]:
        st.subheader(f"📄 {r['label']}")
        st.markdown(r["text"])
        st.divider()
    st.caption(
        "AI can make mistakes. Check important facts in your textbook. "
        f"(Model: {', '.join(sorted(models_used)) or 'none'})"
    )

    st.write("**Was this helpful?**")
    rating = st.feedback("thumbs", key=f"thumbs_{st.session_state['run_id']}")
    if rating is not None:
        st.success(
            "Thanks for your feedback! \U0001F64F"
            if rating == 1
            else "Sorry about that. Please tell us what went wrong in the form below."
        )

    FORM_URL = "https://docs.google.com/forms/d/e/1FAIpQLSd2Ye-U8eMSdyUrBcAU5eMIUvGYhRUPPMwejPOOLcFnH1unIA/viewform"
    st.link_button("\U0001F4AC Give us your feedback (30 seconds)", FORM_URL)
